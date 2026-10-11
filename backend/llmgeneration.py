from langchain.prompts import PromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_ollama import OllamaLLM
from langchain_huggingface import HuggingFaceEmbeddings
from vectorstorage import VectorStorageManager
from langchain_core.output_parsers import JsonOutputParser
from fpdf import FPDF

from docx import Document
from docx.shared import Pt

class QuestionGenerator:
    def __init__(self, ollama_model="gemma3:4b", embed_model="sentence-transformers/all-MiniLM-L12-v2"):
        self.vsm = VectorStorageManager()
        self.embeddings = HuggingFaceEmbeddings(model_name=embed_model)
        self.llm = OllamaLLM(model=ollama_model)

        self.prompts = {
            "multiple": self._build_multiple_choice_prompt(),
            "corta": self._build_short_answer_prompt(),
            "espacio": self._build_fill_blank_prompt(),
            "emparejar": self._build_match_prompt()
        }

    def _build_multiple_choice_prompt(self):
        return PromptTemplate(
            template="""
                You are an assistant for question-answering tasks.

                **LANGUAGE MATCHING INSTRUCTIONS:**
                1. First, identify the language of the input text
                2. Generate ALL content (questions, options, correct answers, explanations) in the SAME language
                3. If the text is in Spanish, respond in Spanish
                4. If the text is in English, respond in English
                5. If the text is in another language, respond in that language
                6. Do NOT translate or mix languages

                Your job is to generate exactly {number} multiple-choice question(s) based on the input text. 
                Each question must include:
                - 4 answer options
                - One correct answer
                - A short explanation of the correct answer

                **You must generate exactly {number} question(s). Do not generate more or fewer.**

                Format the output exactly as a JSON objetct like this (if number is 1, only include one object in the list):

                {{
                    "questions": [
                        {{
                            "question": "Your question here",
                            "options": ["option1", "option2", "option3", "option4"],
                            "correct_answer": "Correct option text",
                            "explanation": "Brief explanation of why this answer is correct"
                        }}
                    ]
                }}

                Text:
                {documents}
                """,
            input_variables=["number", "documents"]
        )

    def _build_short_answer_prompt(self):
        return PromptTemplate(
            template="""
                You are an assistant for question-answering tasks.

                **LANGUAGE MATCHING INSTRUCTIONS:**
                1. First, identify the language of the input text
                2. Generate ALL content (questions, correct answers, explanations) in the SAME language
                3. If the text is in Spanish, respond in Spanish
                4. If the text is in English, respond in English
                5. If the text is in another language, respond in that language
                6. Do NOT translate or mix languages

                Your job is to generate exactly {number} short answer question(s) based on the input text. 
                Each question must include:
                - One correct answer
                - A short explanation of the correct answer

                **Important Instructions:**
                - The correct answer must be **only one word** (no phrases, no multiple words).
                - Do not include punctuation in the answer.
                - The explanation can be longer, but the answer itself must remain a **single word**.

                You must generate **exactly {number}** question(s). Do not generate more or fewer.**

                Format the output exactly as a JSON objetct like this (if number is 1, only include one object in the list):

                {{
                    "questions": [
                        {{
                            "question": "Your question here",
                            "correct_answer": "One word correct answer text",
                            "explanation": "Brief explanation of why this answer is correct"
                        }}
                    ]
                }}

                Text:
                {documents}
                """,
            input_variables=["number", "documents"]
        )

    def _build_fill_blank_prompt(self):
        return PromptTemplate(
            template="""
                You are an assistant for question-generation tasks.

                **LANGUAGE MATCHING INSTRUCTIONS:**
                1. First, identify the language of the input text
                2. Generate ALL content (questions, correct answers, explanations) in the SAME language
                3. If the text is in Spanish, respond in Spanish
                4. If the text is in English, respond in English
                5. If the text is in another language, respond in that language
                6. Do NOT translate or mix languages

                Your job is to generate exactly {number} **fill-in-the-blank** question(s) based on the input text.

                Each question must follow these rules:
                - It must contain a **blank space** represented by three underscores: `___`
                - The correct answer must be **exactly one word**
                - Provide a short explanation of why the answer is correct

                **Important:**
                - The correct answer must be a single word only (no phrases, no hyphenated words).
                - Do not include any punctuation in the answer.
                - The explanation can be longer, but the answer itself must remain one word.

                You must generate **exactly {number}** fill-in-the-blank question(s). Do not generate more or fewer.

                Format the output exactly as a JSON objetct in this format:

                {{
                "questions": [
                    {{
                    "question": "The process of converting sunlight into energy is called ___",
                    "correct_answer": "photosynthesis",
                    "explanation": "Photosynthesis is the process plants use to convert sunlight into chemical energy."
                    }}
                ]
                }}

                Text:
                {documents}
                """,
            input_variables=["number", "documents"]
        )

    def _build_match_prompt(self):
        return PromptTemplate(
            template="""
                You are an assistant for question-answering tasks.
                
                **LANGUAGE MATCHING INSTRUCTIONS:**
                1. First, identify the language of the input text
                2. Generate ALL content (questions, terms, concepts) in the SAME language
                3. If the text is in Spanish, respond in Spanish
                4. If the text is in English, respond in English
                5. If the text is in another language, respond in that language
                6. Do NOT translate or mix languages

                Your job is to generate one matching-type question based on the input text. 
                This question must include:
                - Exactly {number} terms
                - Exactly {number} concepts that correctly match each term

                Ensure the number of terms and concepts is exactly {number}. Do not generate more or fewer.

                Format the output exactly as a JSON objetct in this format:

                {{
                "questions": [
                    {{
                    "terms": ["term1", "term2", "..."],
                    "concepts": ["matching concept1", "matching concept2", "..."]
                    }}
                ]
                }}

                Text:
                {documents}
                """,
            input_variables=["number", "documents"]
        )

    def get_documents(self, namespace: str, filename: str) -> str:
        docs = self.vsm.get_docs(namespace, filename)
        return "\n".join([' '.join(d.split()) for d in docs])

    def generate_questions(self, question_type: str, number: int, namespace: str, filename: str) -> str:
        if question_type not in self.prompts:
            raise ValueError(f"Unsupported question type: {question_type}")

        documents = self.get_documents(namespace, filename)
        prompt = self.prompts[question_type]
        chain = prompt | self.llm | StrOutputParser()

        return chain.invoke({"number": number, "documents": documents})

    def generate_all_types(self, types: list[str], number: int, namespace: str, filename: str) -> dict:
        documents = self.get_documents(namespace, filename)
        results = {}
        parser = JsonOutputParser()

        for question_type in types:
            if question_type not in self.prompts:
                raise ValueError(f"Unsupported question type: {question_type}")

            prompt = self.prompts[question_type]
            chain = prompt | self.llm | parser

            try:
                output = chain.invoke({"number": number, "documents": documents})
                results[question_type] = output
            except Exception as e:
                results[question_type] = f"Error generating {question_type}: {str(e)}"

        return results

class PDF(FPDF):
    def __init__(self):
        super().__init__()
        self.add_font("DejaVu", "", "./font/DejaVuSans.ttf", uni=True)
        self.set_font("DejaVu", "", 11)

    def header(self):
        self.set_font("DejaVu", "", 14)
        self.cell(0, 10, "Preguntas generadas", ln=True, align="C")
        self.ln(5)

    def chapter_title(self, label):
        self.set_font("DejaVu", "", 12)
        self.set_text_color(0, 0, 128)
        self.cell(0, 10, f"--- {label.upper()} ---", ln=True)
        self.ln(2)

    def question_block(self, question_data):
        self.set_font("DejaVu", "", 11)

        if isinstance(question_data, dict):
            question_data = [question_data]

        for item in question_data:
            self.set_text_color(0, 0, 0)
            
            if "terms" in item and "concepts" in item:
                self.multi_cell(0, 8, "Empareja los siguientes términos con sus definiciones:")
                self.ln(3)
                
                self.multi_cell(0, 8, "Términos:")
                for i, term in enumerate(item["terms"], start=1):
                    self.multi_cell(0, 8, f"   {i}. {term}")
                
                self.ln(2)
                
                self.multi_cell(0, 8, "Definiciones:")
                for i, concept in enumerate(item["concepts"], start=1):
                    self.multi_cell(0, 8, f"   {chr(64+i)}. {concept}")
                
            else:
                self.multi_cell(0, 8, f"Pregunta: {item['question']}")

                if "options" in item:
                    for i, opt in enumerate(item["options"], start=1):
                        self.multi_cell(0, 8, f"   {chr(64+i)}. {opt}")

                self.set_text_color(0, 128, 0)
                self.multi_cell(0, 8, f"Respuesta: {item['correct_answer']}")

                self.set_text_color(0, 0, 0)
                self.multi_cell(0, 8, f"Explicación: {item['explanation']}")

            self.ln(5)

def create_questions_pdf(question_outputs: dict, output_file="questions.pdf"):
    pdf = PDF()
    pdf.add_page()

    for qtype, raw_json in question_outputs.items():
        pdf.chapter_title(qtype)

        try:
            if isinstance(raw_json, dict) and "questions" in raw_json:
                questions = raw_json["questions"]
            elif isinstance(raw_json, list):
                questions = raw_json
            else:
                raise ValueError("Invalid format: must be a list or contain 'questions'")
            pdf.question_block(questions)
        except Exception as e:
            pdf.set_font("Arial", "I", 10)
            pdf.set_text_color(255, 0, 0)
            pdf.multi_cell(0, 8, f"Error parsing {qtype}: {str(e)}")
            pdf.set_text_color(0, 0, 0)

    pdf.output(output_file)
    print(f"PDF guardado en: {output_file}")

def create_questions_docx(question_outputs: dict, output_file="questions.docx"):
    doc = Document()
    doc.add_heading("Preguntas generadas", level=1)

    for qtype, raw_json in question_outputs.items():
        doc.add_heading(qtype.upper(), level=2)

        try:
            if isinstance(raw_json, dict) and "questions" in raw_json:
                questions = raw_json["questions"]
            elif isinstance(raw_json, list):
                questions = raw_json
            else:
                raise ValueError("Invalid format: must be a list or contain 'questions'")

            for item in questions:
                if "terms" in item and "concepts" in item:
                    doc.add_paragraph("Empareja los siguientes términos con sus definiciones:")
                    doc.add_paragraph("Términos:", style="List Number")
                    for i, term in enumerate(item["terms"], start=1):
                        doc.add_paragraph(f"{i}. {term}", style="List Bullet")
                    doc.add_paragraph("Definiciones:", style="List Number")
                    for i, concept in enumerate(item["concepts"], start=1):
                        doc.add_paragraph(f"{chr(64+i)}. {concept}", style="List Bullet")
                else:
                    doc.add_paragraph(f"Pregunta: {item['question']}", style="Normal")
                    if "options" in item:
                        for i, opt in enumerate(item["options"], start=1):
                            doc.add_paragraph(f"{chr(64+i)}. {opt}", style="List Bullet")
                    doc.add_paragraph(f"Respuesta: {item['correct_answer']}", style="Intense Quote")
                    doc.add_paragraph(f"Explicación: {item['explanation']}", style="Normal")
                doc.add_paragraph("")

        except Exception as e:
            doc.add_paragraph(f"Error parsing {qtype}: {str(e)}", style="Intense Quote")

    doc.save(output_file)
    print(f"Word document saved as: {output_file}")
if __name__ == "__main__":
    qgen = QuestionGenerator()

    #Solo 1 tipo
    #output = qgen.generate_questions(
    #    question_type="fill",
    #    number=5,
    #    namespace="damian",
    #    filename="philo.pdf"
    #)
    #print(output)


    #Para generar json
    results = qgen.generate_all_types(
        types=["corta","multiple", "espacio", "emparejar"],
        number=2,
        namespace="damian",
        filename="philo.pdf"
    )

    print("diccionario")
    print(results)
    print("//////////////////")
    
    for qtype, output in results.items():
        print(f"\n--- {qtype.upper()} ---\n{output}")

    # para crear pdf
    #create_questions_pdf(results, output_file="generated_questions.pdf")
    create_questions_docx(results, output_file="generated_questions.docx")

