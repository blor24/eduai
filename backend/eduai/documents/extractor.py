from opendocument_text import extract_odf_text
import os
from docx import Document
from pptx import Presentation
import fitz
import pytesseract
from PIL import Image
import io
import pdfplumber


class FileContentExtractor:
    def __init__(self):
        self.ocr_lang = "eng+spa"

    def read_docx(self, filepath):
        doc = Document(filepath)
        return "\n".join([para.text for para in doc.paragraphs])

    def read_txt(self, filepath):
        with open(filepath, 'r', encoding='utf-8') as f:
            return f.read()

    def read_pptx(self, filepath):
        presentation = Presentation(filepath)
        text = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    text.append(shape.text)
        return "\n".join(text)

    def render_pdf_pages_to_images(self, pdf_path, zoom_x=1.25, zoom_y=1.25):
        doc = []
        output_folder = './images' + f"/{os.path.basename(pdf_path)}"
        os.makedirs(output_folder, exist_ok=True)
        pdf_doc = fitz.open(pdf_path)

        for page_number in range(len(pdf_doc)):
            page = pdf_doc.load_page(page_number)
            mat = fitz.Matrix(zoom_x, zoom_y)
            pix = page.get_pixmap(matrix=mat)
            output_file = os.path.join(output_folder, f"page_{page_number + 1}.png")
            pix.save(output_file)

            text = pytesseract.image_to_string(output_file, lang=self.ocr_lang)
            doc.append(text)
            print(f"Saved high-res image: {output_file}")

        pdf_doc.close()
        return "\n".join(doc)

    def extract_images_docx(self, docx_path):
        doc = Document(docx_path)
        output_folder = './images' + f"/{os.path.basename(docx_path)}"
        os.makedirs(output_folder, exist_ok=True)

        ocr_texts = []

        for i, rel in enumerate(doc.part.rels.values()):
            if "image" in rel.target_ref:
                image_data = rel.target_part.blob
                image_path = os.path.join(output_folder, f"image_{i}.jpg")
                with open(image_path, "wb") as img_file:
                    img_file.write(image_data)

                try:
                    img = Image.open(io.BytesIO(image_data))
                    text = pytesseract.image_to_string(img, lang=self.ocr_lang)
                    ocr_texts.append(text)
                except Exception as e:
                    print(f"OCR failed on image {image_path}: {e}")

        return "\n".join(ocr_texts)

    def extract_images_pptx(self, pptx_path):
        prs = Presentation(pptx_path)
        output_folder = './images' + f"/{os.path.basename(pptx_path)}"
        os.makedirs(output_folder, exist_ok=True)

        ocr_texts = []

        for i, slide in enumerate(prs.slides):
            for j, shape in enumerate(slide.shapes):
                if hasattr(shape, "image"):
                    image = shape.image
                    ext = image.ext or "jpg"
                    image_path = os.path.join(output_folder, f"slide_{i}_image_{j}.{ext}")

                    with open(image_path, "wb") as img_file:
                        img_file.write(image.blob)

                    try:
                        img = Image.open(io.BytesIO(image.blob))
                        text = pytesseract.image_to_string(img, lang=self.ocr_lang)
                        ocr_texts.append(text)
                    except Exception as e:
                        print(f"OCR failed on image {image_path}: {e}")

        return "\n".join(ocr_texts)

    def is_pdf_mostly_images(self, filepath, min_text_threshold=50):
        try:
            with pdfplumber.open(filepath) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            return len(text.strip()) < min_text_threshold
        except:
            return True

    def is_docx_mostly_images(self, filepath, min_text_threshold=50):
        doc = Document(filepath)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return len(text.strip()) < min_text_threshold

    def is_pptx_mostly_images(self, filepath, min_text_threshold=50):
        prs = Presentation(filepath)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    text.append(shape.text)
        return len("\n".join(text).strip()) < min_text_threshold

    def get_sections(self, file_path, *, source):
        """Devuelve documentos con metadatos; PDF se indexa por página.

        Para archivos ya indexados antes de esta versión no existen números de página;
        estos metadatos aparecen tras una carga nueva o una reindexación completa.
        """
        from langchain_core.documents import Document
        ext = os.path.splitext(file_path)[1].lower()
        if ext != '.pdf':
            text = self.get_text(file_path)
            return [Document(page_content=text, metadata={'source': source})] if text.strip() else []

        sections = []
        with fitz.open(file_path) as pdf:
            for page_number, page in enumerate(pdf, 1):
                text = page.get_text('text').strip()
                if len(text) < 25:
                    # OCR por página sólo si el PDF es escaneado o casi no tiene texto.
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                    picture = Image.open(io.BytesIO(pixmap.tobytes('png')))
                    ocr = pytesseract.image_to_string(picture, lang=self.ocr_lang).strip()
                    if len(ocr) > len(text):
                        text = ocr
                if text:
                    sections.append(Document(page_content=text,
                                             metadata={'source': source, 'page': page_number}))
        return sections

    def get_text(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return self.render_pdf_pages_to_images(file_path) if self.is_pdf_mostly_images(file_path) else self._safe_pdf_read(file_path)

        elif ext == ".docx":
            return self.extract_images_docx(file_path) if self.is_docx_mostly_images(file_path) else self.read_docx(file_path)

        elif ext == ".pptx":
            return self.extract_images_pptx(file_path) if self.is_pptx_mostly_images(file_path) else self.read_pptx(file_path)

        elif ext in (".txt", ".md"):
            return self.read_txt(file_path)

        elif ext in (".odt", ".fodt"):
            return extract_odf_text(file_path)

        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _safe_pdf_read(self, file_path):
        try:
            with pdfplumber.open(file_path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except:
            return self.render_pdf_pages_to_images(file_path)

# EDUAI_FUENTES_V1
